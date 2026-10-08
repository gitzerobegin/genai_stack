#!/usr/bin/env python3
"""Apply Stage A' verifier corrections (V1/V2 verification_log.md section 3) to the raw
"What changed" rows and write the CP1 table (markdown + xlsx).

Usage: python3 -I tools/build_cp1_what_changed.py <repo_root>
"""
import json, os, sys
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

root = sys.argv[1]
os.chdir(root)
rows = json.load(open("work/cp1/what_changed_raw.json", encoding="utf-8"))
rows = [r for r in rows if not r["original"].startswith("Plan / brief")]

# key: exact original label -> dict(append=..., flag=..., replace=(old,new), src=...)
C = {
 # ---- V1 (streams A1-A4) ----
 "OpenRouter – multi-provider": dict(flag="Acquired (pending)", append="Stripe and OpenRouter announced on 19 August 2026 that Stripe will acquire OpenRouter; no completion notice by 8 October 2026; price undisclosed (press: US$7–8bn+). The US$113M Series B (26 May 2026) is now Verified.", src="V1-S059, V1-S060, V1-S061"),
 "Hugging Face – models & APIs": dict(append="TGI GitHub repository archived (read-only) on 21 March 2026.", src="V1-S054, V1-S055"),
 "Postgres – + pgvector": dict(append="0.8.7 released 1 October 2026 (announced 5 October); fixes CVE-2026-103484 (CVSS 8.8); extension licence is the PostgreSQL License.", src="V1-S020, V1-S021, V1-S022"),
 "Mistral OCR – OCR": dict(append="OCR 4.1 released July 2026, GA 26 August 2026 (Mistral pages give 16 July, 26 July and 13 August).", src="V1-S014"),
 "Promptfoo – red-teaming": dict(flag="Acquired (announced; closing not published)", append="OpenAI announced the acquisition on 9 March 2026; no closing notice exists, only the README's 'part of OpenAI' wording.", src="V1-S006"),
 "Phoenix – Atrace": dict(append="Dynatrace completed the Arize acquisition on 1 October 2026 (Dynatrace IR).", src="V1-S005"),
 "Arize – RAG metrics": dict(append="Dynatrace completion date 1 October 2026 (Dynatrace IR).", src="V1-S005"),
 "Tavily – search API": dict(append="Closed 19 February 2026 (Nebius FY2025 Form 20-F). US$275M is a Bloomberg figure, not disclosed by Nebius.", src="V1-S041"),
 "CrewAI – multi-agent": dict(append="'AMP, formerly CrewAI Enterprise' is inferred, not confirmed by a rename notice: earlier materials call it CrewAI Enterprise.", src="V1-S075"),
 "Mem0 – memory layer": dict(append="Open-source v2 replaced graph stores with built-in entity linking (not a queryable graph); deletions recorded under SDK v2.0.1.", src="V1-S043"),
 "Voyage AI – Voyage-3": dict(append="rerank-3 / rerank-3-lite are listed as Preview on MongoDB's model lifecycle page.", src="V1-S024"),
 "Pinecone – managed": dict(append="BYOC announced GA late September 2026; requires Enterprise; some features (Assistant, Inference, on-demand indexes) not yet in BYOC as of late August.", src="V1-S034, V1-S029"),
 "NVIDIA – Embed": dict(append="nemotron-3-embed-1b added in Embedding NIM 2.2; 2.3 is current.", src="V1-S094"),
 "Docling – doc parser": dict(append="Donated by IBM to LF AI & Data in March 2025 (Incubation); Graduate in August 2026.", src="V1-S091"),
 "Unstructured – ETL for docs": dict(append="Adds CMMC 2.0 Level 2. No FedRAMP authorisation found: do not repeat the 'FedRAMP adherence' wording.", src="V1-S090"),
 "Agent Skills – reusable skills": dict(flag="No change (open format; no neutral governance body)", append="Open-standard date (18 December 2025) and independent adoption (OpenAI Codex) confirmed; no neutral governance body; not an AAIF-hosted project on available evidence.", src="V1-S040, V1-S046, V1-S047"),
 "Jina AI – Embeddings v3": dict(append="Elastic completed the acquisition on 9 October 2025.", src="V1-S025"),
 # ---- V2 (streams A5-A8) ----
 'Grok (no version)': dict(replace=("xAI merged into SpaceX (2 February 2026) and was rebranded **SpaceXAI** (July 2026) (Reported).", "SpaceX acquired xAI in an all-stock deal announced and closed on 2 February 2026 (Verified, xAI's own post). The AI unit was rebranded SpaceXAI in mid-2026 (date unverified: May or July). On 4 October 2026 Musk announced a further rename to 'SpaceXSI', not yet effected."), src="V2-S011"),
 'DeepSeek "V4"': dict(append="The V4-Pro phase-out planned for 14 September 2026 was reversed (pricing-page note, 17 September): V4-Pro API continues.", src="V2-S013"),
 'Qwen "3.8"': dict(append="Hosted Qwen3.8-Max launched 3 August 2026; the 2.4T-A95B weights (12 August) are under a custom 'Qwen3.8-Max License', not Apache 2.0 (Reported).", src="V2-S014"),
 '"QI4" (Z logo) "Q4"': dict(replace=("GLM-5.3 and GLM-5.3-Flash (August 2026)", "GLM-5.3 (API mid-August 2026: 14 August per NIST CAISI, 18 August per press; weights c. 28 August under a bespoke licence) and GLM-5.3-Flash (26 August 2026, MIT)"), src="V2-S021, V2-S016"),
 'Claude "Opus 5.5"': dict(append="Also: Claude Fable 5 access was suspended 12 June 2026 and restored 1 July 2026; a related N.D. Cal. ruling is dated 27 August 2026 (see L1-anthropic record). Disclosure: the author is an Anthropic model.", src="V2-S004, V2-S068"),
 'Meta "Llama (new: Muse)"': dict(append="Meta Model API GA confirmed (Connect, 23–24 September 2026).", src="V2-S019"),
 'Mistral "Medium 3.1"': dict(append="Large 4 preview is API-only; weights targeted for 27 October 2026 (VentureBeat).", src="V2-S017"),
 "(not in graphic) Kong AI Gateway": dict(append="AI Gateway 2.2 GA 30 September 2026.", src="V2-S034"),
 "(not in graphic) Guardrails AI": dict(append="Cutoff date conflict: repository notice says 25 August 2026, docs migration guide says 6 August 2026.", src="V2-S071, V2-S072"),
 "(not in graphic) Microsoft Entra Agent ID": dict(replace=("GA in 2026 (What's new page dated 1 May 2026). Security features require Microsoft Agent 365 licences.", "GA in April 2026 (Entra release log). Security features need Agent 365, included in M365 E7 and sold as an add-on to E5/A5/Business Premium."), src="V2-S032"),
 "(not in graphic) Okta / Auth0 for AI Agents": dict(append="GA-date conflict for Okta Agent SSO resolved in favour of 24 August 2026.", src="V2-S035"),
 "(not in graphic) Cedar / AgentCore Policy": dict(append="AgentCore Policy GA in 13 Regions; policies now authored in Dogwood, an open-source superset of Cedar.", src="V2-S033"),
 "(not in graphic) OAuth 2.1 / MCP authorisation": dict(append="2026-07-28 also removes SSE resumability and deprecates the Roots, Sampling and Logging features.", src="V2-S031"),
 "(not in graphic) LiteLLM": dict(append="v1.83.0 uploaded to PyPI 31 March 2026 05:08 UTC (30 March US time); exposure window about 40 minutes (LiteLLM) to about 3 hours (Snyk).", src="V2-S027, V2-S028"),
 "(not in graphic) Google Apigee AI gateway": dict(append="Vertex AI → Gemini Enterprise Agent Platform rename dated April 2026 (22 April release notes).", src="V2-S037"),
 "(none – C6) FinOps FOCUS": dict(flag="No change (token column deferred)", replace=("**FOCUS 1.5 is scoped for AI model identity and input/output tokens.**", "FOCUS 1.4 ratified 4 June 2026. FOCUS 1.5 (no ratification date) adds AI pricing dimensions on SkuPriceDetails and four model properties (ModelDeveloper, ModelFamily, ModelId, ModelVersion); a first-class input/output token-type column is deferred."), src="V2-S046"),
 "(none – C5) Prompts-as-code pattern": dict(append="OpenAI's Promptfoo acquisition was announced 9 March 2026; closing date not published.", src="V2-S042"),
 "(none – C5) Langfuse Prompt Management": dict(append="ClickHouse announced the acquisition on 16 January 2026.", src="V2-S041"),
 "(none – C7) Lakera": dict(append="Completion date 22 October 2025 (Check Point Q3 2025 results); consideration about US$201.8m (20-F, second hand).", src="V2-S038"),
 "(none – C7) HashiCorp Vault": dict(append="2.1.1 is the latest version in the changelog, but a v2.1.2 git tag exists: check before citing.", src="V2-S061"),
 "(none – C7) HiddenLayer": dict(append="Series B dated 2 September 2026.", src="V2-S047"),
 "EBA outsourcing guidelines (2019)": dict(replace=("with a two-year transition", "reference EBA/GL/2026/09; application date not yet fixed (translations pending); critical or important arrangements to be reviewed within two years of application"), src="V2-S054"),
 "SEC predictive data analytics rule": dict(replace=("Withdrawn 17 June 2025", "Withdrawn by Commission action on 12 June 2025; notice published 17 June 2025"), src="V2-S058"),
 "ISO/IEC 42001 (42005/42006 \"if published\")": dict(append="ISO/IEC 42006 published 7 July 2025.", src="V2-S057"),
 "OWASP Top 10 for LLM Applications 2025": dict(append="Date conflict stands (press release 2 September 2026 vs archive 3 August 2026): cite as August–September 2026.", src="V2-S056"),
}
applied = set()
for r in rows:
    c = C.get(r["original"])
    r["verification"] = "Checked by A′; no change"
    if not c:
        continue
    applied.add(r["original"])
    if "replace" in c:
        old, new = c["replace"]
        if old in r["current"]:
            r["current"] = r["current"].replace(old, new)
        else:
            r["current"] += " **A′ correction:** " + new
    if "append" in c:
        r["current"] += " **A′:** " + c["append"]
    if "flag" in c:
        r["flag"] = c["flag"]
    r["sources"] += ", " + c["src"]
    r["verification"] = "Corrected/extended by A′ (" + c["src"] + ")"
missing = set(C) - applied
if missing:
    print("WARNING unmatched corrections:", missing)

LAYER = {"A1": "L9 / L8", "A2": "L7 / L6", "A3": "L5 / L4", "A4": "L3 / L2", "A5": "L1", "A6": "C1–C4", "A7": "C5–C8", "A8": "Regulation"}
os.makedirs("checkpoints/CP1", exist_ok=True)
json.dump(rows, open("checkpoints/CP1/what_changed.json", "w", encoding="utf-8"), indent=2, ensure_ascii=False)

wb = Workbook(); ws = wb.active; ws.title = "what_changed"
ws.append(["#", "Area", "Original label / assumption", "Current reality (as of 8 October 2026)", "Flag", "Sources", "Stage A′ verification"])
for i, r in enumerate(rows, 1):
    ws.append([i, LAYER[r["stream"]], r["original"], r["current"].replace("**", ""), r["flag"], r["sources"], r["verification"]])
for c in ws[1]:
    c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="1F3864")
for i, w in enumerate([5, 11, 32, 90, 26, 30, 30], 1):
    ws.column_dimensions[get_column_letter(i)].width = w
for row in ws.iter_rows():
    for c in row:
        c.alignment = Alignment(wrap_text=True, vertical="top")
ws.freeze_panes = "D2"; ws.auto_filter.ref = ws.dimensions
os.makedirs("Enterprise_GenAI_Stack_Oct2026/05_Data", exist_ok=True)
wb.save("Enterprise_GenAI_Stack_Oct2026/05_Data/what_changed.xlsx")

with open("checkpoints/CP1/02_What_Changed_Since_Original_Diagram.md", "w", encoding="utf-8") as f:
    f.write("# What changed since the original diagram\n\n")
    f.write("| | |\n|---|---|\n| **As of** | 8 October 2026 |\n| **Basis** | Stage A research (8 streams), corrected by Stage A′ adversarial verification (V1, V2) |\n| **Rows** | %d: 80 graphic tiles, the control-plane candidates (not in the graphic) and the plan's regulatory assumptions |\n| **Also in** | `Enterprise_GenAI_Stack_Oct2026/05_Data/what_changed.xlsx` (filterable) |\n\n" % len(rows))
    f.write("Source IDs resolve in `Enterprise_GenAI_Stack_Oct2026/06_References/bibliography.xlsx`. Text marked **A′** was added or corrected by the verifiers.\n\n")
    cur = None
    for i, r in enumerate(rows, 1):
        if r["stream"] != cur:
            cur = r["stream"]
            f.write("\n## %s\n\n| # | Original | Current reality | Flag | Sources |\n|---:|---|---|---|---|\n" % LAYER[cur])
        f.write("| %d | %s | %s | %s | %s |\n" % (i, r["original"], r["current"], r["flag"], r["sources"]))
print(len(rows), "rows;", len(applied), "corrections applied")
