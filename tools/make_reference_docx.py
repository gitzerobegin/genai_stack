#!/usr/bin/env python3
"""Brand the pandoc Word template (tools/templates/reference.docx) in the Veyan style.

Usage: python3 -I tools/make_reference_docx.py <repo_root>
Starts from tools/templates/reference_base.docx (the unbranded template) every time, so it can be re-run.
Adds the Veyan lockup (visual icon + word mark) to the page header, the V-and-eye mark with the edition line to the
footer, and Veyan colours to the title and heading styles. Every .docx built with --reference-doc inherits these.
"""
import os, sys
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from edition import E, PKG
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

root = sys.argv[1]; os.chdir(root)
NAVY, GOLD, BLUE, GREY = RGBColor(0x0B, 0x1B, 0x33), RGBColor(0xD4, 0xA1, 0x3A), RGBColor(0x2E, 0x6D, 0xA4), RGBColor(0x6B, 0x7A, 0x8F)
doc = Document("tools/templates/reference_base.docx")

for st in doc.styles:
    n = (st.name or "").lower()
    if n in ("title", "subtitle", "heading 1", "heading 2", "heading 3", "heading 4", "toc heading"):
        st.font.color.rgb = NAVY if n != "subtitle" else BLUE
        if n in ("title", "heading 1", "heading 2"): st.font.name = "Arial"; st.font.bold = True

def bottom_rule(par, color="D4A13A", sz="12"):
    pPr = par._p.get_or_add_pPr(); b = OxmlElement("w:pBdr"); e = OxmlElement("w:bottom")
    for k, v in (("w:val", "single"), ("w:sz", sz), ("w:space", "4"), ("w:color", color)): e.set(qn(k), v)
    b.append(e); pPr.append(b)

sec = doc.sections[0]
sec.header.is_linked_to_previous = False; sec.footer.is_linked_to_previous = False
for p in list(sec.header.paragraphs) + list(sec.footer.paragraphs): p._p.getparent().remove(p._p)
hp = sec.header.add_paragraph(); hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
hp.add_run().add_picture("brand/veyan_lockup.png", width=Inches(1.45))
bottom_rule(hp)
fp = sec.footer.add_paragraph(); fp.alignment = WD_ALIGN_PARAGRAPH.LEFT
fp.add_run().add_picture("brand/veyan_mark.png", height=Inches(0.2))
r = fp.add_run("   Veyan · The Enterprise GenAI Stack · " + E["view_label"][0].lower() + E["view_label"][1:] + " · not a description of any firm's platform")
r.font.size = Pt(7.5); r.font.color.rgb = GREY
doc.save("tools/templates/reference.docx")
print("branded tools/templates/reference.docx")
