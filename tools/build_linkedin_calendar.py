#!/usr/bin/env python3
"""Build the LinkedIn content calendar (XLSX) and the standalone series document (DOCX) from
work/stageC2/linkedin_series.md.

Usage: python3 -I tools/build_linkedin_calendar.py <repo_root> [--start YYYY-MM-DD]
Without --start, dates are left blank and a 'Suggested date' formula column computes them from a single
start-date cell you fill in (the date of Post 1). No day of the week is fixed: the first post of each week is the
stack post, the second the control post, two days apart by default; overwrite any date freely.
Writes Enterprise_GenAI_Stack_Oct2026/07_LinkedIn/Content_Calendar.xlsx and LinkedIn_Series.docx
"""
import json, os, re, subprocess, sys
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

root = sys.argv[1]; os.chdir(root)
OUT = "Enterprise_GenAI_Stack_Oct2026/07_LinkedIn"; os.makedirs(OUT, exist_ok=True)
text = open("work/stageC2/linkedin_series.md", encoding="utf-8").read()
posts = re.split(r"(?m)^(?=### Post \d+)", text)[1:]

def sub(block, name):
    m = re.search(r"(?ms)^#### %s\s*\n(.*?)(?=^#### |^### |^## |\Z)" % re.escape(name), block)
    return m.group(1).strip() if m else ""
def field(block, name):
    m = re.search(r"\*\*%s:\*\*\s*(.+)" % re.escape(name), block)
    return m.group(1).strip() if m else ""

WX = json.load(open("work/stageC2/worked_example_build.json", encoding="utf-8")) if os.path.exists("work/stageC2/worked_example_build.json") else {"steps": []}
WXS = {s["post"]: s for s in WX["steps"]}
rows = []
for b in posts:
    head = b.split("\n", 1)[0]
    m = re.match(r"### Post (\d+)\s*·\s*Week (\d+)\s*·\s*(.+)", head)
    if not m: continue
    n, wk, theme = int(m.group(1)), int(m.group(2)), m.group(3).strip()
    day = "Introduction" if n == 0 else ("Stack post" if n % 2 else "Control post")
    full = sub(b, "Full post")
    hook = next((l.strip() for l in full.split("\n") if l.strip()), "")
    rows.append(dict(n=n, week=wk, day=day, theme=theme, pair=field(b, "Pair"), tension=field(b, "Tension"), hook=hook[:300],
                     words=len(re.sub(r"\[[^\]]*\]", "", full).split()), visual=sub(b, "Suggested visual")[:500],
                     tags=sub(b, "Hashtags").replace("\n", " ")[:80], reverify=sub(b, "Re-verify before posting").replace("\n", " ")[:500]))

wb = Workbook(); ws = wb.active; ws.title = "calendar"
ws["A1"] = "Series start date (the date you publish Post 1):"; ws["A1"].font = Font(bold=True)
ws["E1"] = None; ws["E1"].fill = PatternFill("solid", fgColor="FFF2CC")
ws["F1"] = "← enter the date of Post 1; suggested dates below compute from it (any weekday; edit freely)"
start_arg = next((sys.argv[i + 1] for i, a in enumerate(sys.argv) if a == "--start"), None)
if start_arg:
    import datetime; ws["E1"] = datetime.date.fromisoformat(start_arg); ws["E1"].number_format = "dd mmm yyyy"
cols = ["#", "Week", "Slot", "Suggested date", "Time (UK)", "Theme", "Worked example step", "Pair", "Tension", "Hook (first line)", "Words", "Hashtags", "Visual brief",
        "Re-verify before posting", "Status", "Cleared", "Published URL", "Responses / notes", "Full post and visual"]
ws.append([]); ws.append(cols)
for c in ws[3]:
    c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="1F3864"); c.alignment = Alignment(wrap_text=True, vertical="top")
for r in sorted(rows, key=lambda x: x["n"]):
    i = ws.max_row + 1
    off = -3 if r["n"] == 0 else (r["week"] - 1) * 7 + (0 if r["day"] == "Stack post" else 2)
    wx = WXS.get(r["n"], {})
    ws.append([r["n"], r["week"], r["day"], '=IF($E$1="","",$E$1+%d)' % off, "08:00", r["theme"], (wx.get("stage", "") + ": " + wx.get("adds", "")) if wx else "", r["pair"], r["tension"], r["hook"], r["words"],
               r["tags"], r["visual"], r["reverify"], "Draft", "Not required (CP4b)", "", "",
               "07_LinkedIn/LinkedIn_Series.docx, Post %d · visual 08_Graphic/linkedin/P%02d.png" % (r["n"], r["n"])])
    ws.cell(i, 4).number_format = "dd mmm yyyy"
widths = [5, 6, 9, 15, 9, 28, 45, 30, 40, 45, 7, 22, 45, 45, 11, 16, 25, 30, 30]
for i, w in enumerate(widths, 1): ws.column_dimensions[get_column_letter(i)].width = w
for row in ws.iter_rows(min_row=4):
    for c in row: c.alignment = Alignment(wrap_text=True, vertical="top")
dv = DataValidation(type="list", formula1='"Draft,Anecdote added,Scheduled,Published,Skipped"', allow_blank=True)
ws.add_data_validation(dv); dv.add("O4:O%d" % max(ws.max_row, 4))
ws.freeze_panes = "G4"
wb.save(os.path.join(OUT, "Content_Calendar.xlsx"))
# Word edition: each post's rendered visual sits under its "Suggested visual" brief (08_Graphic/linkedin/P<NN>.png)
VDIR = "Enterprise_GenAI_Stack_Oct2026/08_Graphic/linkedin"
def add_visual(block):
    m = re.match(r"### Post (\d+)", block)
    if not m: return block
    pid = "P%02d" % int(m.group(1)); png = os.path.join(VDIR, pid + ".png")
    if not os.path.exists(png): return block
    fig = "\n![Visual for Post %d](%s){width=4.2in}\n\n*Visual: `08_Graphic/linkedin/%s.png` (1080 × 1350, LinkedIn portrait); editable source `%s.md`, also as .html and .pdf.*\n\n" % (int(m.group(1)), png, pid, pid)
    return re.sub(r"(?ms)(^#### Suggested visual\s*\n.*?)(?=^#### )", lambda x: x.group(1).rstrip("\n") + "\n" + fig, block, count=1)
pieces = re.split(r"(?m)^(?=### )", text)
os.makedirs("work/stageD", exist_ok=True)
wx_tab = "" if not WX["steps"] else ("\n\n## Appendix: the worked example, step by step\n\n" + WX.get("use_case", "") + "\n\n| Step | Stage | Adds | Now works | Never |\n|---|---|---|---|---|\n" +
          "".join("| %d | %s | %s | %s | %s |\n" % (s["post"], s["stage"], s["adds"], s["now_works"], s["boundary"]) for s in WX["steps"]))
open("work/stageD/LinkedIn_Series.pandoc.md", "w", encoding="utf-8").write("".join(add_visual(b) for b in pieces) + wx_tab)
r = subprocess.run(["pandoc", "work/stageD/LinkedIn_Series.pandoc.md", "-f", "markdown-tex_math_dollars-raw_tex-implicit_figures", "-o", os.path.join(OUT, "LinkedIn_Series.docx"),
                    "--toc", "--toc-depth=2", "--reference-doc=tools/templates/reference.docx"], capture_output=True, text=True)
print(len(rows), "posts; pandoc", r.returncode)
