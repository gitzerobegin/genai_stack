#!/usr/bin/env python3
"""Merge Stage A stream outputs into the product dataset and source bibliography.

Usage: python3 -I tools/build_dataset.py <repo_root>

Reads   work/stageA/*/products.json, work/stageA/*/sources.csv, work/stageA/A8_Regulation/regulatory_facts.json
Writes  Enterprise_GenAI_Stack_Oct2026/05_Data/products.json
        Enterprise_GenAI_Stack_Oct2026/05_Data/products.xlsx   (one row per product, fact cells flattened)
        Enterprise_GenAI_Stack_Oct2026/05_Data/regulatory_facts.json
        Enterprise_GenAI_Stack_Oct2026/06_References/bibliography.xlsx (sources + claim map)
        work/stageA/_integrity_report.md
"""
import csv, glob, json, os, sys
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

root = sys.argv[1]
os.chdir(root)
PKG = "Enterprise_GenAI_Stack_Oct2026"
LAYER_ORDER = ["L9", "L8", "L7", "L6", "L5", "L4", "L3", "L2", "L1", "C1", "C2", "C3", "C4", "C5", "C6", "C7", "C8"]

def is_fact(x):
    return isinstance(x, dict) and "v" in x and "label" in x

def flat(x):
    if x is None:
        return ""
    if is_fact(x):
        v = x["v"]
        v = "; ".join(map(str, v)) if isinstance(v, list) else str(v)
        return v
    if isinstance(x, dict):
        return "; ".join("%s: %s" % (k, flat(v)) for k, v in x.items())
    if isinstance(x, list):
        return "; ".join(flat(i) for i in x)
    return str(x)

def walk_facts(obj, path=""):
    if is_fact(obj):
        yield path, obj
    elif isinstance(obj, dict):
        for k, v in obj.items():
            yield from walk_facts(v, (path + "." if path else "") + k)

products, sources, issues = [], {}, []
for d in sorted(glob.glob("work/stageA/A*_*/")) + sorted(glob.glob("work/stageA_verify/V*/")):
    stream = os.path.basename(d.rstrip("/")).split("_")[0]
    sp = os.path.join(d, "sources.csv")
    if os.path.exists(sp):
        with open(sp, newline="", encoding="utf-8") as f:
            for row in csv.DictReader(f):
                row = {k.strip(): (v or "").strip() for k, v in row.items() if k}
                if row.get("id"):
                    if row["id"] in sources:
                        issues.append("duplicate source id %s" % row["id"])
                    row["stream"] = stream
                    sources[row["id"]] = row
    elif "stageA_verify" not in d:
        issues.append("%s: no sources.csv" % d)
    pp = os.path.join(d, "products.json")
    if os.path.exists(pp):
        try:
            arr = json.load(open(pp, encoding="utf-8"))
        except Exception as e:
            issues.append("%s: invalid JSON (%s)" % (pp, e)); continue
        for p in arr:
            p["_stream"] = stream
            products.append(p)

# integrity: every src id resolves; Verified/Reported cells carry sources
ids = set()
for p in products:
    if p.get("id") in ids:
        issues.append("duplicate product id %s" % p.get("id"))
    ids.add(p.get("id"))
    for path, cell in walk_facts(p):
        for s in cell.get("src") or []:
            if s not in sources:
                issues.append("%s.%s cites unknown source %s" % (p.get("id"), path, s))
        if cell.get("label") in ("Verified fact", "Reported") and not cell.get("src"):
            issues.append("%s.%s labelled %s without source" % (p.get("id"), path, cell.get("label")))

def sort_key(p):
    l = p.get("layer", "Z")
    return (LAYER_ORDER.index(l) if l in LAYER_ORDER else 99, p.get("id", ""))
products.sort(key=sort_key)

os.makedirs(PKG + "/05_Data", exist_ok=True)
os.makedirs(PKG + "/06_References", exist_ok=True)
json.dump(products, open(PKG + "/05_Data/products.json", "w", encoding="utf-8"), indent=2, ensure_ascii=False)
reg = []
rp = "work/stageA/A8_Regulation/regulatory_facts.json"
if os.path.exists(rp):
    reg = json.load(open(rp, encoding="utf-8"))
    json.dump(reg, open(PKG + "/05_Data/regulatory_facts.json", "w", encoding="utf-8"), indent=2, ensure_ascii=False)
    for r in reg:
        for path, cell in walk_facts(r):
            for s in cell.get("src") or []:
                if s not in sources:
                    issues.append("%s.%s cites unknown source %s" % (r.get("id"), path, s))

HDR = PatternFill("solid", fgColor="1F3864")
def style(ws, widths):
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF"); c.fill = HDR
        c.alignment = Alignment(wrap_text=True, vertical="top")
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = ws.dimensions
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")

COLS = ["id", "layer", "original_label", "current_name", "company", "category", "version_or_lineup", "licence_model",
        "strategic_direction", "status_events", "what_it_does", "stack_position", "dependencies", "integration_model",
        "deployment", "certifications", "gdpr_residency", "security_features", "access_controls", "enterprise_support",
        "pricing", "infra_cost_note", "ecosystem", "adoption_signals", "maturity", "strategic_risk",
        "assessment", "classification", "scores", "sources", "stage_a_notes", "last_verified"]
wb = Workbook(); ws = wb.active; ws.title = "products"
ws.append(COLS + ["labels_used", "npv_fields"])
for p in products:
    labels = sorted({c.get("label", "") for _, c in walk_facts(p)})
    npv = [path for path, c in walk_facts(p) if c.get("label") == "Not publicly verified"]
    ws.append([flat(p.get(c)) for c in COLS] + ["; ".join(labels), "; ".join(npv)])
style(ws, [18, 6, 24] + [32] * (len(COLS) - 3) + [24, 30])

ws2 = wb.create_sheet("fact_cells")
ws2.append(["product_id", "field", "value", "label", "confidence", "source_ids"])
for p in products:
    for path, c in walk_facts(p):
        ws2.append([p.get("id"), path, flat(c), c.get("label"), c.get("conf", ""), "; ".join(c.get("src") or [])])
style(ws2, [20, 28, 70, 20, 10, 24])
wb.save(PKG + "/05_Data/products.xlsx")

# bibliography: sources + claim map
SCOLS = ["id", "stream", "url", "title", "publisher", "source_type", "accessed", "access_status", "archive_path", "used_for"]
bw = Workbook(); b1 = bw.active; b1.title = "sources"
b1.append(SCOLS)
for sid in sorted(sources):
    b1.append([sources[sid].get(c, "") for c in SCOLS])
style(b1, [12, 7, 50, 40, 22, 20, 12, 22, 45, 40])
b2 = bw.create_sheet("claim_map")
b2.append(["record_id", "field", "claim", "label", "confidence", "source_id", "url"])
for rec in products + reg:
    for path, c in walk_facts(rec):
        for s in c.get("src") or []:
            b2.append([rec.get("id"), path, flat(c)[:500], c.get("label"), c.get("conf", ""), s, sources.get(s, {}).get("url", "")])
style(b2, [20, 28, 70, 18, 10, 12, 50])
bw.save(PKG + "/06_References/bibliography.xlsx")

with open("work/stageA/_integrity_report.md", "w", encoding="utf-8") as f:
    f.write("# Stage A integrity report\n\nProducts: %d · Regulatory records: %d · Sources: %d\n\n" % (len(products), len(reg), len(sources)))
    from collections import Counter
    f.write("## Products per layer\n\n" + "\n".join("- %s: %d" % kv for kv in sorted(Counter(p.get("layer") for p in products).items(), key=lambda kv: LAYER_ORDER.index(kv[0]) if kv[0] in LAYER_ORDER else 99)) + "\n\n")
    lab = Counter(c.get("label") for p in products for _, c in walk_facts(p))
    f.write("## Fact-cell labels\n\n" + "\n".join("- %s: %d" % kv for kv in lab.most_common()) + "\n\n")
    st = Counter(s.get("access_status", "").split(" ")[0] for s in sources.values())
    f.write("## Source access status\n\n" + "\n".join("- %s: %d" % kv for kv in st.most_common()) + "\n\n")
    ty = Counter(s.get("source_type", "") for s in sources.values())
    f.write("## Source types\n\n" + "\n".join("- %s: %d" % kv for kv in ty.most_common()) + "\n\n")
    f.write("## Issues (%d)\n\n" % len(issues) + "\n".join("- " + i for i in issues[:500]) + "\n")
print(json.dumps({"products": len(products), "regulatory": len(reg), "sources": len(sources), "issues": len(issues)}))
